import streamlit as st
import pandas as pd
import os
import glob
from openpyxl.utils import get_column_letter
from datetime import datetime
import json
import tempfile
import io

# =====================================================================================
# FUNÇÃO UNIFICADA: PROCESSA O ARQUIVO INDEPENDENTE DO TIPO
# =====================================================================================
def processar_arquivo_generico(caminho_csv):
    """
    Lê um arquivo CSV e tenta identificar automaticamente se é PRODUTO ACABADO
    ou BOBINA baseando-se na estrutura dos dados.
    """
    try:
        try:
            df = pd.read_csv(caminho_csv, header=None, encoding='utf-8', dtype=str)
        except UnicodeDecodeError:
            df = pd.read_csv(caminho_csv, header=None, encoding='latin1', dtype=str)
    except Exception as e:
        return None, f"Erro ao ler arquivo: {e}"

    dados_processados = []
    
    if df.empty:
        return pd.DataFrame(), None

    for index, row in df.iterrows():
        if len(row) < 4:
            continue

        col_zero = str(row[0]).strip()
        
        # Filtro de sujeira (ignora o cabeçalho "DATE,..." ou linhas vazias)
        if "date" in col_zero.lower() or not col_zero or not col_zero[0].isdigit():
            continue

        # ====================================================================
        # DETECTOR AUTOMÁTICO DE FORMATO DO CSV (CELULARES/APPS DIFERENTES)
        # ====================================================================
        # Se a primeira coluna tem espaço e é grande (ex: "2026-06-21 11:40:14 -0300")
        if len(col_zero) > 10 and " " in col_zero:
            # FORMATO NOVO (Ex: N6F7)
            partes_dt = col_zero.split(" ")
            dt_leitura_original = partes_dt[0] # Pega só o '2026-06-21'
            hr_leitura = partes_dt[1] if len(partes_dt) > 1 else ""
            
            # O dado lido (CODECONTENT) está na coluna 4 (índice 3)
            dados_lidos = str(row[3]).strip() if len(row) > 3 else ""
            coluna_tipo = str(row[4]).strip() if len(row) > 4 else ""
            
            # Converte a data de YYYY-MM-DD para DD/MM/YYYY
            try:
                obj_date = datetime.strptime(dt_leitura_original, '%Y-%m-%d')
                dt_formatada = obj_date.strftime('%d/%m/%Y')
            except:
                dt_formatada = dt_leitura_original

        else:
            # FORMATO ANTIGO (Ex: N6F4)
            dt_leitura_original = col_zero
            hr_leitura = str(row[1]).strip() if len(row) > 1 else ""
            
            # O dado lido (CODECONTENT) está na coluna 5 (índice 4)
            coluna_tipo = str(row[3]).strip() if len(row) > 3 else ""
            dados_lidos = str(row[4]).strip() if len(row) > 4 else ""

            # Converte a data de MM-DD-YYYY para DD/MM/YYYY
            try:
                obj_date = datetime.strptime(dt_leitura_original, '%m-%d-%Y')
                dt_formatada = obj_date.strftime('%d/%m/%Y')
            except:
                dt_formatada = dt_leitura_original


        # Estrutura padrão da linha
        nova_linha = {
            "Data da Leitura": dt_formatada,
            "Hora da Leitura": hr_leitura,
            "Filial": None,
            "Código": None,
            "Armazém": None,
            "Lote": None,
            "Peso": None,
            "Localização": os.path.splitext(os.path.basename(caminho_csv))[0]
        }

        # ====================================================================
        # TESTE 1: É PRODUTO ACABADO CLÁSSICO? (Padrão: "XXX-XXX - YYY")
        # ====================================================================
        processado_como_pa = False
        if " -" in dados_lidos:
            try:
                partes_maiores = dados_lidos.split(" -", 1)
                
                parte_esq = partes_maiores[0].split("-")
                filial = parte_esq[0].strip() if len(parte_esq) > 0 else ""
                codigo = parte_esq[1].strip() if len(parte_esq) > 1 else ""
                
                parte_dir = partes_maiores[1].split("-") if len(partes_maiores) > 1 else []
                
                armazem = parte_dir[0].strip() if len(parte_dir) > 0 else ""
                lote = parte_dir[1].strip() if len(parte_dir) > 1 else ""
                peso_str = parte_dir[2].strip() if len(parte_dir) > 2 else "0"
                
                try:
                    peso_val = float(peso_str) / 1000.0
                except:
                    peso_val = None

                nova_linha["Filial"] = filial
                nova_linha["Código"] = codigo
                nova_linha["Armazém"] = armazem
                nova_linha["Lote"] = lote
                nova_linha["Peso"] = peso_val
                
                dados_processados.append(nova_linha)
                processado_como_pa = True
            except Exception:
                pass
        
        if processado_como_pa:
            continue

        # ====================================================================
        # TESTE 2: É BOBINA? (Ou Bobina com estrutura completa)
        # ====================================================================
        
        lote_b = "erro"
        peso_b = None 

        # CASO 1: CODE128 (Asteriscos)
        if coluna_tipo == 'Code128' or '*' in dados_lidos:
            if ' ' in dados_lidos:
                 lote_b, peso_b = "erro de leitura", None
            elif '*' in dados_lidos:
                try:
                    partes = dados_lidos.split('*')
                    if dados_lidos.startswith('*'): 
                        lote_b = partes[3].strip()
                        peso_b = float(partes[2].strip()) / 1000.0
                    else: 
                        lote_b = partes[2].strip()
                        peso_b = float(partes[1].strip()) / 1000.0
                except:
                    lote_b, peso_b = "erro Code128/*", None
            elif dados_lidos.isdigit() and len(dados_lidos) <= 5:
                 try:
                    peso_b, lote_b = float(dados_lidos)/1000.0, ""
                 except:
                    peso_b, lote_b = None, dados_lidos
            else:
                 lote_b, peso_b = dados_lidos, None

        # CASO 2: QR CODE / DATAMATRIX
        elif coluna_tipo in ['QR_CODE', 'QR', 'CODE_39', 'CODE_128'] or '{' in dados_lidos or ',' in dados_lidos:
            
            # 2.1 JSON
            if '{' in dados_lidos and '}' in dados_lidos:
                try:
                    partes = dados_lidos.split('{', 1)
                    identificador = partes[0].strip('"-')
                    dados_json = json.loads('{' + partes[1])
                    peso_b = float(dados_json.get('peso', 0))
                    lote_b = identificador
                except:
                    lote_b = "erro QR/JSON"

            # 2.2 NOVO FORMATO COM VÍRGULA
            elif ',' in dados_lidos and '-' in dados_lidos:
                try:
                    partes_virgula = dados_lidos.split(',')
                    
                    if len(partes_virgula) > 1 and partes_virgula[-1].replace('.', '', 1).isdigit():
                        peso_decimal = partes_virgula[-1].strip()
                        
                        parte_lote_completa = ','.join(partes_virgula[:-1])
                        partes_hifen = parte_lote_completa.split('-')
                        
                        # Extração de Filial/Código/Armazém se disponível
                        if len(partes_hifen) >= 4:
                            nova_linha["Filial"] = partes_hifen[0].strip()
                            nova_linha["Código"] = partes_hifen[1].strip()
                            nova_linha["Armazém"] = partes_hifen[2].strip()

                        lote_b = partes_hifen[-2].strip()
                        
                        peso_completo_str = f"{partes_hifen[-1].strip()},{peso_decimal}"
                        peso_b = float(peso_completo_str.replace(',', '.'))
                        
                    else:
                        raise ValueError("Formato virgula invalido")
                except:
                    try:
                        partes = dados_lidos.split('-')
                        lote_b = partes[3].strip()
                        peso_b = float(partes[-1].strip()) / 1000.0
                    except:
                        lote_b = "erro QR/FormatoVirgula"
            
            # 2.3 Formato Antigo (Só hifens)
            else:
                 try:
                     partes = dados_lidos.split('-')
                     if len(partes) >= 4:
                        lote_b = partes[3].strip() 
                        peso_b = float(partes[-1].strip()) / 1000.0
                     else:
                        lote_b = dados_lidos
                        peso_b = None
                 except:
                     lote_b = dados_lidos
                     peso_b = None

        else:
            lote_b = dados_lidos
            peso_b = None 
        
        nova_linha["Lote"] = lote_b
        nova_linha["Peso"] = peso_b
        dados_processados.append(nova_linha)

    return pd.DataFrame(dados_processados), None

# =====================================================================================
# LEITOR DE MÃO (Tomate MDK 103S): FUNÇÕES DE PROCESSAMENTO
# =====================================================================================
PREFIXO_LOC = "LOC:"
COR_DUPLICADO = "FFF6CC"   # amarelo
COR_SEM_LOC = "FFE8CC"     # laranja
COR_ERRO = "FFD9D9"        # vermelho

def interpretar_leitura_mao(texto):
    """
    Recebe o texto de UMA leitura e devolve um dicionário com o tipo:
      - {"tipo": "local", "local": "N4F1"}
      - {"tipo": "item", "Filial":..., "Código":..., "Armazém":..., "Lote":..., "Peso":...}
      - {"tipo": "erro", "motivo": "..."}
    Padrão do item: FILIAL-PRODUTO -ARMAZÉM-LOTE-PESO
    """
    texto = str(texto).strip()

    # É uma localização? (QR com prefixo LOC:)
    if texto.upper().startswith(PREFIXO_LOC):
        local = texto[len(PREFIXO_LOC):].strip()
        if local:
            return {"tipo": "local", "local": local}
        return {"tipo": "erro", "motivo": "QR de localização sem nome"}

    # É um item? Precisa ter o " -" que separa o produto do armazém
    if " -" not in texto:
        return {"tipo": "erro", "motivo": f'Fora do padrão: "{texto}"'}

    lado_esq, lado_dir = texto.split(" -", 1)
    partes_esq = lado_esq.split("-", 1)
    partes_dir = lado_dir.split("-")

    if len(partes_esq) < 2 or len(partes_dir) < 3:
        return {"tipo": "erro", "motivo": f'Fora do padrão: "{texto}"'}

    filial = partes_esq[0].strip()
    codigo = partes_esq[1].strip()
    armazem = partes_dir[0].strip()
    lote = "-".join(partes_dir[1:-1]).strip()   # tudo entre o armazém e o peso
    peso_txt = partes_dir[-1].strip().replace(",", ".")

    try:
        peso = float(peso_txt) / 1000.0
    except ValueError:
        return {"tipo": "erro", "motivo": f'Peso inválido: "{texto}"'}

    if not all([filial, codigo, armazem, lote]):
        return {"tipo": "erro", "motivo": f'Campo vazio: "{texto}"'}

    return {"tipo": "item", "Filial": filial, "Código": codigo, "Armazém": armazem,
            "Lote": lote, "Peso": peso}


def processar_planilha_leitor(arquivo):
    """
    Lê a planilha gerada pelo leitor de mão (todas as abas, célula por célula,
    de cima para baixo) e devolve um DataFrame com os itens já separados.
    Cada aba/arquivo começa SEM localização (não herda do anterior).
    """
    abas = pd.read_excel(arquivo, sheet_name=None, header=None, dtype=str)
    linhas = []

    for _, df in abas.items():
        local_atual = None
        for _, row in df.iterrows():
            for valor in row:
                if pd.isna(valor) or not str(valor).strip():
                    continue
                r = interpretar_leitura_mao(valor)

                if r["tipo"] == "local":
                    local_atual = r["local"]
                    continue

                if r["tipo"] == "erro":
                    linhas.append({"Filial": "", "Código": "", "Armazém": "", "Lote": "",
                                   "Peso": None, "Localização": local_atual or "SEM LOCALIZAÇÃO",
                                   "Observação": r["motivo"], "_status": "erro"})
                    continue

                linhas.append({"Filial": r["Filial"], "Código": r["Código"],
                               "Armazém": r["Armazém"], "Lote": r["Lote"], "Peso": r["Peso"],
                               "Localização": local_atual or "SEM LOCALIZAÇÃO",
                               "Observação": "" if local_atual else "Lido antes de qualquer QR de localização",
                               "_status": "ok" if local_atual else "semloc"})

    return pd.DataFrame(linhas, columns=["Filial", "Código", "Armazém", "Lote", "Peso",
                                         "Localização", "Observação", "_status"])


def marcar_duplicados(df):
    """Marca os lotes lidos mais de uma vez (todas as ocorrências)."""
    df = df.copy()
    df["_dup"] = False
    itens = df["_status"] != "erro"
    contagem = df.loc[itens, "Lote"].value_counts()
    repetidos = contagem[contagem > 1]

    for lote, qtd in repetidos.items():
        mask = itens & (df["Lote"] == lote)
        df.loc[mask, "_dup"] = True
        aviso = f"Lote duplicado ({qtd} leituras)"
        df.loc[mask, "Observação"] = df.loc[mask, "Observação"].apply(
            lambda o: f"{o} | {aviso}" if o else aviso)
    return df


def cor_da_linha(status, dup):
    """Prioridade: erro (vermelho) > sem localização (laranja) > duplicado (amarelo)."""
    if status == "erro":
        return COR_ERRO
    if status == "semloc":
        return COR_SEM_LOC
    if dup:
        return COR_DUPLICADO
    return None


def resumo_por_local(df):
    """Quantidade de itens por localização, na ordem em que foram lidos."""
    itens = df[df["_status"] != "erro"]
    if itens.empty:
        return pd.DataFrame(columns=["Localização", "Itens"])
    ordem = list(dict.fromkeys(itens["Localização"]))
    if "SEM LOCALIZAÇÃO" in ordem:
        ordem.remove("SEM LOCALIZAÇÃO")
        ordem.insert(0, "SEM LOCALIZAÇÃO")
    qtd = itens["Localização"].value_counts()
    return pd.DataFrame({"Localização": ordem, "Itens": [int(qtd[l]) for l in ordem]})


def gerar_excel_leitor(df, df_locais, data_inventario="", comparacao=None):
    """Monta o Excel final: aba 'Inventario Geral' (com cores) + aba 'Itens por Localização' + aba 'Info'."""
    from openpyxl.styles import PatternFill, Font

    colunas = ["Filial", "Código", "Descrição", "Armazém", "Lote", "Peso", "Localização", "Observação"]
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df[colunas].to_excel(writer, index=False, sheet_name="Inventario Geral")
        ws = writer.sheets["Inventario Geral"]

        # Cabeçalho em negrito + congelado
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"

        for i, (status, dup) in enumerate(zip(df["_status"], df["_dup"]), start=2):
            # Colunas de texto ficam como TEXTO (mantém zeros à esquerda); Peso com 3 casas
            for col, nome in enumerate(colunas, start=1):
                ws.cell(row=i, column=col).number_format = "0.000" if nome == "Peso" else "@"

            cor = cor_da_linha(status, dup)
            if cor:
                fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
                for col in range(1, len(colunas) + 1):
                    ws.cell(row=i, column=col).fill = fill

        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

        # Aba 2: itens por localização
        df_locais.to_excel(writer, index=False, sheet_name="Itens por Localização")
        ws2 = writer.sheets["Itens por Localização"]
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        ws2.column_dimensions["A"].width = 24
        ws2.column_dimensions["B"].width = 10
        for i, local in enumerate(df_locais["Localização"], start=2):
            if local == "SEM LOCALIZAÇÃO":
                fill = PatternFill(start_color=COR_SEM_LOC, end_color=COR_SEM_LOC, fill_type="solid")
                ws2.cell(row=i, column=1).fill = fill
                ws2.cell(row=i, column=2).fill = fill
        total = len(df_locais) + 2
        ws2.cell(row=total, column=1, value="Total").font = Font(bold=True)
        ws2.cell(row=total, column=2, value=int(df_locais["Itens"].sum())).font = Font(bold=True)

        # Aba 3: informações do inventário (usada pela aba Ordem de Carregamento)
        pd.DataFrame({"Data do inventário": [data_inventario],
                      "Gerado em": [agora_br().strftime("%d/%m/%Y %H:%M")]}
                     ).to_excel(writer, index=False, sheet_name="Info")
        ws3 = writer.sheets["Info"]
        ws3.column_dimensions["A"].width = 20
        ws3.column_dimensions["B"].width = 20

        # Abas 4 e 5: comparação com o estoque (só se a opção foi marcada)
        if comparacao:
            for nome_aba, chave, cor in (("Falta", "falta", COR_ERRO), ("Sobra", "sobra", "DBEAFE")):
                tabela = comparacao[chave]
                tabela.to_excel(writer, index=False, sheet_name=nome_aba)
                wsx = writer.sheets[nome_aba]
                for cell in wsx[1]:
                    cell.font = Font(bold=True)
                wsx.freeze_panes = "A2"
                fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
                for i in range(2, len(tabela) + 2):
                    for j, nome_col in enumerate(tabela.columns, start=1):
                        c = wsx.cell(row=i, column=j)
                        c.fill = fill
                        c.number_format = "0.000" if "(t)" in nome_col else "@"
                for col in wsx.columns:
                    larg = max((len(str(c.value)) for c in col if c.value is not None), default=0)
                    wsx.column_dimensions[get_column_letter(col[0].column)].width = min(larg + 4, 60)

    output.seek(0)
    return output


# =====================================================================================
# ETIQUETAS DE LOCALIZAÇÃO: QR CODE + PDF
# =====================================================================================
LAYOUTS_ETIQUETA = {
    "1 por folha (A4 inteira)": (1, 1),
    "2 por folha (1 x 2)": (1, 2),
    "6 por folha (2 x 3)": (2, 3),
    "12 por folha (3 x 4)": (3, 4),
}

def gerar_qr_png(texto, escala=10):
    """Gera o QR Code em PNG (bytes)."""
    import segno
    buf = io.BytesIO()
    segno.make_qr(texto, error="m").save(buf, kind="png", scale=escala, border=4)
    buf.seek(0)
    return buf


def limpar_lista_locais(texto):
    """Uma localização por linha, sem vazios, sem repetidos, sem o prefixo LOC: se alguém digitou."""
    locais = []
    for linha in texto.splitlines():
        nome = linha.strip()
        if nome.upper().startswith(PREFIXO_LOC):
            nome = nome[len(PREFIXO_LOC):].strip()
        if nome and nome not in locais:
            locais.append(nome)
    return locais


def gerar_pdf_etiquetas(locais, colunas, linhas):
    """Monta o PDF A4 com as etiquetas (QR com 'LOC:' + nome grande embaixo)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    largura, altura = A4
    margem = 12 * 2.835  # 12 mm em pontos
    cel_w = (largura - 2 * margem) / colunas
    cel_h = (altura - 2 * margem) / linhas
    por_pagina = colunas * linhas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle("Etiquetas de Localização")

    for idx, local in enumerate(locais):
        pos = idx % por_pagina
        if idx > 0 and pos == 0:
            c.showPage()
        col = pos % colunas
        lin = pos // colunas
        x = margem + col * cel_w
        y = altura - margem - (lin + 1) * cel_h

        # Borda tracejada para recorte
        c.setDash(4, 3)
        c.setStrokeColorRGB(0.6, 0.6, 0.6)
        c.rect(x + 3, y + 3, cel_w - 6, cel_h - 6)
        c.setDash()

        # Tamanho do texto proporcional à etiqueta
        tam_fonte = max(14, min(cel_w, cel_h) * 0.16)
        espaco_texto = tam_fonte * 1.9

        # QR o maior possível dentro da etiqueta
        lado_qr = min(cel_w * 0.80, cel_h - espaco_texto - 20)
        qr_x = x + (cel_w - lado_qr) / 2
        qr_y = y + espaco_texto + (cel_h - espaco_texto - lado_qr) / 2
        c.drawImage(ImageReader(gerar_qr_png(PREFIXO_LOC + local)), qr_x, qr_y, lado_qr, lado_qr)

        # Nome da localização embaixo (diminui a fonte se o nome for comprido)
        tam_nome = tam_fonte
        while c.stringWidth(local, "Helvetica-Bold", tam_nome) > cel_w * 0.9 and tam_nome > 8:
            tam_nome -= 1
        c.setFont("Helvetica-Bold", tam_nome)
        c.setFillColorRGB(0.15, 0.15, 0.2)
        c.drawCentredString(x + cel_w / 2, y + tam_fonte * 0.9, local)

    c.save()
    buf.seek(0)
    return buf

# =====================================================================================
# DESCRIÇÃO DOS PRODUTOS (aba "Dados_Produtos" atualizada pelo robô)
# =====================================================================================
ID_PLANILHA_SISTEMA = "1jODOp_SJUKWp1UaSmW_xJgkkyqDUexa56_P5QScAv3s"
ABA_PRODUTOS = "Dados_Produtos"

@st.cache_data(ttl=6 * 60 * 60, show_spinner=False)
def carregar_produtos():
    """
    Lê a lista código -> descrição da planilha do robô.
    Fica guardada na memória por 6 horas: só a 1ª conversão nesse período consulta o Google.
    Se der erro, NÃO guarda (tenta de novo na próxima conversão).
    """
    import gspread
    try:
        gc = gspread.service_account_from_dict(dict(st.secrets["gcp_service_account"]))
    except Exception:
        gc = gspread.service_account(filename="credentials.json")   # quando roda no seu PC
    valores = gc.open_by_key(ID_PLANILHA_SISTEMA).worksheet(ABA_PRODUTOS).get_all_values()
    return {str(l[0]).strip(): str(l[1]).strip() for l in valores[1:] if len(l) >= 2 and str(l[0]).strip()}

# =====================================================================================
# COMPARAÇÃO COM O ESTOQUE DO PROTHEUS (aba "Dados_Estoque" atualizada pelo robô)
# =====================================================================================
@st.cache_data(ttl=10 * 60, show_spinner=False)
def carregar_estoque():
    """
    Lê o estoque (Dados_Estoque) e a hora em que o robô o atualizou (Status_Robo!B2),
    numa única chamada ao Google. Fica guardado por 10 minutos.
    """
    import gspread
    try:
        gc = gspread.service_account_from_dict(dict(st.secrets["gcp_service_account"]))
    except Exception:
        gc = gspread.service_account(filename="credentials.json")   # quando roda no seu PC
    blocos = gc.open_by_key(ID_PLANILHA_SISTEMA).values_batch_get(["Dados_Estoque", "Status_Robo!B2"])
    faixas = blocos.get("valueRanges", [])
    valores = faixas[0].get("values", []) if faixas else []
    hora = ""
    if len(faixas) > 1 and faixas[1].get("values"):
        hora = str(faixas[1]["values"][0][0])
    if not valores:
        return pd.DataFrame(), hora
    cab = valores[0]
    linhas = [l + [""] * (len(cab) - len(l)) for l in valores[1:]]
    return pd.DataFrame(linhas, columns=cab), hora


def formatar_ton(v):
    """Mostra o peso com 3 casas e vírgula (ex: 1,338). Vazio se não tiver valor."""
    return "" if pd.isna(v) else f"{v:.3f}".replace(".", ",")


def comparar_com_estoque(df_mao, estoque):
    """
    Cruza os lotes lidos com o estoque do Protheus da(s) mesma(s) filial(is) das etiquetas.
      - Falta: está no Protheus e não foi lido
      - Sobra: foi lido e não está no Protheus
    """
    itens = df_mao[df_mao["_status"] != "erro"]
    filiais = sorted(set(f for f in itens["Filial"] if f))
    prefixos = tuple(f"{f}-" for f in filiais)

    est = estoque.copy()
    for col in ("FILIAL", "LOTE", "COD", "PRODUTO", "ARMAZEM", "QTDE", "DIAS.ESTOQUE"):
        if col not in est.columns:
            est[col] = ""
        est[col] = est[col].astype(str).str.strip()
    est = est[est["FILIAL"].str.startswith(prefixos) & (est["LOTE"] != "")]
    est = est.drop_duplicates(subset=["LOTE", "ARMAZEM"])

    lotes_protheus = set(est["LOTE"])
    lotes_lidos = set(l for l in itens["Lote"] if l)
    encontrados = lotes_lidos & lotes_protheus
    falta = lotes_protheus - lotes_lidos
    sobra = lotes_lidos - lotes_protheus

    df_falta = est[est["LOTE"].isin(falta)].copy()
    df_falta["QTDE"] = pd.to_numeric(df_falta["QTDE"].str.replace(",", "."), errors="coerce")
    df_falta["FILIAL"] = df_falta["FILIAL"].str[:2]          # "05-PINHEIRAL" -> "05"
    df_falta = df_falta.rename(columns={"FILIAL": "Filial", "LOTE": "Lote", "COD": "Código",
                                        "PRODUTO": "Descrição", "ARMAZEM": "Armazém",
                                        "QTDE": "Saldo Protheus (t)", "DIAS.ESTOQUE": "Entrada no estoque"})
    df_falta = df_falta[["Filial", "Armazém", "Lote", "Código", "Descrição", "Saldo Protheus (t)",
                         "Entrada no estoque"]].sort_values(["Filial", "Armazém", "Lote"])

    df_sobra = itens[itens["Lote"].isin(sobra)]
    df_sobra = (df_sobra.groupby("Lote", sort=False)
                .agg({"Filial": "first", "Armazém": "first", "Código": "first", "Descrição": "first",
                      "Peso": "first", "Localização": lambda x: " / ".join(dict.fromkeys(x))})
                .reset_index()[["Filial", "Armazém", "Lote", "Código", "Descrição", "Peso", "Localização"]]
                .rename(columns={"Peso": "Peso lido (t)"}))

    return {
        "filiais": filiais,
        "n_protheus": len(lotes_protheus), "n_lidos": len(lotes_lidos),
        "n_encontrados": len(encontrados), "n_falta": len(falta), "n_sobra": len(sobra),
        "acuracidade": (len(encontrados) / len(lotes_protheus) * 100) if lotes_protheus else None,
        "falta": df_falta.reset_index(drop=True), "sobra": df_sobra,
    }

# =====================================================================================
# BASE DE LOCALIZAÇÃO "VIVA" (aba "Base_Localizacao" na planilha Sistema Dox)
# =====================================================================================
ABA_BASE_LOC = "Base_Localizacao"
CABECALHO_BASE_LOC = ["DATA_HORA", "FILIAL", "ARMAZEM", "LOTE", "CODIGO", "LOCALIZACAO", "ORIGEM"]
DIAS_LEITURA_ANTIGA = 30   # no roteiro, leituras mais velhas que isso ficam em amarelo


def agora_br():
    """Data/hora de Brasília (o servidor do Streamlit Cloud fica em outro fuso)."""
    from datetime import timezone, timedelta
    return datetime.now(timezone(timedelta(hours=-3))).replace(tzinfo=None)


def _cliente_google():
    """Conexão com o Google Sheets (Secrets no site; credentials.json quando roda no PC)."""
    import gspread
    try:
        return gspread.service_account_from_dict(dict(st.secrets["gcp_service_account"]))
    except Exception:
        return gspread.service_account(filename="credentials.json")


def linhas_para_base(df_mao, origem):
    """Monta as linhas a gravar: só itens válidos e COM localização."""
    agora = agora_br().strftime("%d/%m/%Y %H:%M:%S")
    ok = df_mao[(df_mao["_status"] != "erro") & (df_mao["Localização"] != "SEM LOCALIZAÇÃO")]
    return [[agora, r["Filial"], r["Armazém"], r["Lote"], r["Código"], r["Localização"], origem]
            for _, r in ok.iterrows()]


def gravar_base_localizacao(linhas):
    """Acrescenta as linhas no final da aba, numa única chamada. Tenta de novo se der 429."""
    import time
    from gspread.exceptions import APIError
    if not linhas:
        return 0
    aba = _cliente_google().open_by_key(ID_PLANILHA_SISTEMA).worksheet(ABA_BASE_LOC)
    for tentativa in range(3):
        try:
            aba.append_rows(linhas, value_input_option="RAW")
            carregar_base_localizacao.clear()   # a próxima consulta já vê as leituras novas
            return len(linhas)
        except APIError as e:
            if ("429" in str(e) or "Quota" in str(e)) and tentativa < 2:
                time.sleep(20)
                continue
            raise


def montar_base(valores, hoje=None):
    """
    Recebe as linhas da aba (com cabeçalho) e devolve:
      - base:  {lote: [localizações da leitura MAIS RECENTE]}
      - lidos: {lote: ("06/10 · há 2 dias", dias)}
    """
    hoje = hoje or agora_br()
    mais_recente = {}   # lote -> (data, [locais])
    for linha in valores[1:]:
        if len(linha) < 6:
            continue
        data_txt, lote, local = linha[0].strip(), linha[3].strip(), linha[5].strip()
        if not lote or not local:
            continue
        try:
            data = datetime.strptime(data_txt, "%d/%m/%Y %H:%M:%S")
        except ValueError:
            continue
        atual = mais_recente.get(lote)
        if atual is None or data > atual[0]:
            mais_recente[lote] = (data, [local])
        elif data == atual[0] and local not in atual[1]:
            atual[1].append(local)            # mesma leitura em 2 locais

    base, lidos = {}, {}
    for lote, (data, locais) in mais_recente.items():
        base[lote] = locais
        dias = (hoje.date() - data.date()).days
        quando = "hoje" if dias <= 0 else ("há 1 dia" if dias == 1 else f"há {dias} dias")
        lidos[lote] = (f"{data.strftime('%d/%m')} · {quando}", dias)
    return base, lidos


@st.cache_data(ttl=10 * 60, show_spinner=False)
def carregar_base_localizacao():
    """Lê a Base de Localização inteira (1 chamada). Fica guardada por 10 minutos."""
    valores = _cliente_google().open_by_key(ID_PLANILHA_SISTEMA).worksheet(ABA_BASE_LOC).get_all_values()
    return valores

# =====================================================================================
# ORDEM DE CARREGAMENTO: FUNÇÕES
# =====================================================================================
def ler_base_inventario(arquivos):
    """
    Lê um ou mais Excel de inventário (gerados pelo Leitor de Mão ou pelo App Celular)
    e devolve:
      - base: {lote: [localizações]}
      - datas: lista de datas do inventário encontradas (aba 'Info' do Leitor de Mão)
    """
    base = {}
    datas = []
    for arq in arquivos:
        abas = pd.read_excel(arq, sheet_name=None, dtype=str)
        df = abas.get("Inventario Geral")
        if df is None or "Lote" not in df.columns or "Localização" not in df.columns:
            raise ValueError(f'O arquivo "{arq.name}" não tem a aba "Inventario Geral" com as colunas Lote e Localização.')

        for lote, local in zip(df["Lote"], df["Localização"]):
            if pd.isna(lote) or pd.isna(local):
                continue
            lote, local = str(lote).strip(), str(local).strip()
            if not lote or not local or local == "SEM LOCALIZAÇÃO":
                continue
            base.setdefault(lote, [])
            if local not in base[lote]:
                base[lote].append(local)

        info = abas.get("Info")
        if info is not None and "Data do inventário" in info.columns and not info.empty:
            datas.append(str(info["Data do inventário"].iloc[0]).strip())
    return base, datas


def buscar_local(base, lote):
    """Procura o lote na base (também tenta com zeros à esquerda, caso algum tenha se perdido)."""
    if lote in base:
        return base[lote]
    if lote.isdigit():
        for chave in (lote.lstrip("0"), lote.zfill(11)):
            if chave in base:
                return base[chave]
    return []


def ler_ordem_carregamento(arquivo):
    """
    Lê o PDF 'Pick-list de Carregamento' do Protheus e devolve:
      - cab: {'precarga', 'carga'}
      - itens: lista com seq, descrição, qtd, lote, cliente e a posição da coluna 'Localizac'
    """
    import re
    import pdfplumber
    from collections import defaultdict

    itens, cab = [], {}
    with pdfplumber.open(arquivo) as pdf:
        for num_pag, pag in enumerate(pdf.pages):
            palavras = pag.extract_words()
            linhas = defaultdict(list)
            for w in palavras:
                linhas[round(w["top"])].append(w)

            col_loc = [w for w in palavras if w["text"].startswith("Localiza")]
            col_cli = [w for w in palavras if w["text"] == "Cliente"]
            col_mun = [w for w in palavras if w["text"] == "Municipio"]
            topos = sorted(linhas)

            for i, topo in enumerate(topos):
                lw = sorted(linhas[topo], key=lambda w: w["x0"])
                textos = [w["text"] for w in lw]

                # Cabeçalho: a linha de baixo do "Pre-Car" tem os números
                if any(t.startswith("Pre-Car") for t in textos) and i + 1 < len(topos):
                    valores = sorted(linhas[topos[i + 1]], key=lambda w: w["x0"])
                    if len(valores) >= 2:
                        cab = {"precarga": valores[0]["text"], "carga": valores[1]["text"]}

                # Linha de item: começa com Seq de 3 dígitos e tem um lote de 11 dígitos
                if not col_loc or not re.fullmatch(r"\d{3}", textos[0]):
                    continue
                x0_loc, x1_loc = col_loc[0]["x0"], col_loc[0]["x1"]
                lotes = [w for w in lw if re.fullmatch(r"\d{11}", w["text"]) and w["x0"] > x1_loc]
                if not lotes:
                    continue
                antes = [w for w in lw if w["x0"] < x0_loc]   # Seq, Descrição, Qtd Embarca, Qtd Peças
                cliente = ""
                if col_cli and col_mun:
                    cliente = " ".join(w["text"] for w in lw
                                       if col_cli[0]["x0"] - 2 <= w["x0"] < col_mun[0]["x0"] - 2)
                itens.append({
                    "pagina": num_pag, "seq": textos[0],
                    "desc": " ".join(w["text"] for w in antes[1:-2]),
                    "qtd": antes[-2]["text"] if len(antes) >= 3 else "",
                    "lote": lotes[0]["text"], "cliente": cliente,
                    "x0": x0_loc, "x1": x1_loc, "top": lw[0]["top"], "bottom": lw[0]["bottom"],
                })
    return cab, itens


def carimbar_ordem(arquivo, itens, base):
    """Escreve a localização por cima do '_________' da coluna Localizac. Sem localização = fica em branco."""
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas
    from collections import defaultdict

    arquivo.seek(0)
    leitor = PdfReader(arquivo)
    saida = PdfWriter()
    por_pagina = defaultdict(list)
    for it in itens:
        por_pagina[it["pagina"]].append(it)

    for num_pag, pagina in enumerate(leitor.pages):
        larg, alt = float(pagina.mediabox.width), float(pagina.mediabox.height)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(larg, alt))
        escreveu = False
        for it in por_pagina.get(num_pag, []):
            locais = buscar_local(base, it["lote"])
            if not locais:
                continue
            escreveu = True
            texto = " / ".join(locais)
            y = alt - it["bottom"]
            largura_col = it["x1"] - it["x0"]
            # Tampa o "_________" com um retângulo branco
            c.setFillColorRGB(1, 1, 1)
            c.rect(it["x0"] - 2, y - 1.5, largura_col + 4, it["bottom"] - it["top"] + 3, stroke=0, fill=1)
            # Diminui a fonte se o texto não couber
            tam = 8.5
            while c.stringWidth(texto, "Helvetica-Bold", tam) > largura_col + 8 and tam > 5:
                tam -= 0.5
            c.setFillColorRGB(0, 0, 0)
            c.setFont("Helvetica-Bold", tam)
            c.drawString(it["x0"] - 1, y + 0.5, texto)
        c.save()
        if escreveu:
            buf.seek(0)
            pagina.merge_page(PdfReader(buf).pages[0])
        saida.add_page(pagina)
    return saida


def gerar_roteiro(cab, itens, base, origem_txt, lidos=None):
    """Folha 'Roteiro de Separação': itens agrupados por localização, sem localização no final."""
    from pypdf import PdfReader
    from collections import defaultdict
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm)
    estilos = getSampleStyleSheet()
    elementos = []

    n_ok = sum(1 for it in itens if buscar_local(base, it["lote"]))
    elementos.append(Paragraph(
        f"<b>ROTEIRO DE SEPARAÇÃO</b> &nbsp;&nbsp; Carga {cab.get('carga', '')} &nbsp;·&nbsp; "
        f"Pré-carga {cab.get('precarga', '')}", estilos["Title"]))
    elementos.append(Paragraph(
        f"{len(itens)} itens · {n_ok} localizados · {len(itens) - n_ok} sem localização "
        f"&nbsp;&nbsp;|&nbsp;&nbsp; <i>Localizações conforme {origem_txt}</i>"
        + (f" &nbsp;|&nbsp; <font backColor='#FFF6CC'>&nbsp;amarelo&nbsp;</font> = lido há mais de "
           f"{DIAS_LEITURA_ANTIGA} dias" if lidos is not None else ""), estilos["Normal"]))
    elementos.append(Spacer(1, 6))

    SEM = "~SEM"  # chave que fica por último na ordenação
    grupos = defaultdict(list)
    for it in itens:
        locais = buscar_local(base, it["lote"])
        grupos[locais[0] if locais else SEM].append((it, locais))

    if lidos is not None:
        dados = [["Localização", "Seq.", "Lote", "Descrição", "Qtd (t)", "Cliente", "Lido em", "OK"]]
        larguras = [38, 12, 26, 74, 16, 50, 36, 12]
    else:
        dados = [["Localização", "Seq.", "Lote", "Descrição", "Qtd (t)", "Cliente", "OK"]]
        larguras = [38, 14, 30, 80, 18, 60, 12]
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    linha = 1
    for chave in sorted(grupos):
        inicio = linha
        for k, (it, locais) in enumerate(grupos[chave]):
            rotulo = ("SEM LOCALIZAÇÃO" if chave == SEM else chave) if k == 0 else ""
            desc = it["desc"]
            if len(locais) > 1:
                desc += f"  (também em {' / '.join(locais[1:])})"
            if lidos is not None:
                txt, dias = lidos.get(it["lote"], ("", 0)) if locais else ("", 0)
                dados.append([rotulo, it["seq"], it["lote"], desc, it["qtd"], it["cliente"], txt, ""])
                if dias > DIAS_LEITURA_ANTIGA:
                    estilo.append(("BACKGROUND", (6, linha), (6, linha), colors.HexColor("#FFF6CC")))
                    estilo.append(("FONTNAME", (6, linha), (6, linha), "Helvetica-Bold"))
            else:
                dados.append([rotulo, it["seq"], it["lote"], desc, it["qtd"], it["cliente"], ""])
            linha += 1
        if chave == SEM:
            estilo.append(("BACKGROUND", (0, inicio), (-1, linha - 1), colors.HexColor("#FFE8CC")))
        estilo.append(("LINEBELOW", (0, linha - 1), (-1, linha - 1), 1.2, colors.black))
        estilo.append(("FONTNAME", (0, inicio), (0, linha - 1), "Helvetica-Bold"))

    tabela = Table(dados, colWidths=[l * mm for l in larguras], repeatRows=1)
    tabela.setStyle(TableStyle(estilo))
    elementos.append(tabela)
    doc.build(elementos)
    buf.seek(0)
    return PdfReader(buf)


# =====================================================================================
# INTERFACE DO STREAMLIT (UI)
# =====================================================================================

st.set_page_config(page_title="Conversor de Inventário Dox", layout="wide")
st.title("Conversor de Inventário Unificado")
st.markdown("---")

aba_celular, aba_mao, aba_etiquetas, aba_ordem = st.tabs(
    ["📱 App Celular", "🔫 Leitor de Mão", "🏷️ Etiquetas de Localização", "🚚 Ordem de Carregamento"])

# =====================================================================================
# ABA 1: APP CELULAR (sem alterações)
# =====================================================================================
with aba_celular:

    # --- LÓGICA DO BOTÃO LIMPAR (SESSION STATE) ---
    if 'uploader_key' not in st.session_state:
        st.session_state.uploader_key = 0

    def limpar_lista():
        st.session_state.uploader_key += 1

    col1, col2 = st.columns([2, 1])

    with col1:
        # O parametro 'key' muda cada vez que clicamos em Limpar, resetando o widget
        uploaded_files = st.file_uploader(
            "Importar arquivos .csv (Aceita Bobina e Produto Acabado misturados)",
            type="csv",
            accept_multiple_files=True,
            key=str(st.session_state.uploader_key)
        )
        
        # --- FRASE DE CONTAGEM E BOTÃO LIMPAR ---
        if uploaded_files:
            qtd_arquivos = len(uploaded_files)
            col_msg, col_btn = st.columns([3, 1])
            with col_msg:
                st.info(f"📂 **{qtd_arquivos} arquivos anexados.**")
            with col_btn:
                st.button("🗑️ Limpar Lista", on_click=limpar_lista, type="secondary", use_container_width=True)

    with col2:
        st.info("Configurações de Saída")
        nome_arquivo_usuario = st.text_input("Nome do Arquivo Final (sem .xlsx):", value="")

    # --- BOTÃO CONVERTER ---
    st.markdown("###") # Espaçamento
    if st.button("Converter Arquivos", type="primary"):
        if not uploaded_files:
            st.warning("⚠️ Por favor, carregue pelo menos um arquivo .csv.")
        else:
            with st.spinner("Processando..."):
                try:
                    todos_dfs = []
                    for uploaded_file in uploaded_files:
                        with tempfile.NamedTemporaryFile(delete=False, suffix='.csv') as tmp_file:
                            tmp_file.write(uploaded_file.getbuffer())
                            tmp_path = tmp_file.name
                        
                        df_temp, erro = processar_arquivo_generico(tmp_path)
                        
                        if erro:
                            st.error(f"Erro no arquivo {uploaded_file.name}: {erro}")
                        elif not df_temp.empty:
                            df_temp["Localização"] = uploaded_file.name.replace('.csv', '')
                            todos_dfs.append(df_temp)
                        
                        os.unlink(tmp_path)

                    if todos_dfs:
                        df_final = pd.concat(todos_dfs, ignore_index=True)
                        
                        # Tratamento estético
                        if "Armazém" in df_final.columns:
                            df_final["Armazém"] = df_final["Armazém"].fillna('').apply(lambda x: str(x).split('.')[0].zfill(2) if str(x).replace('.','').isdigit() else str(x))
                        
                        output = io.BytesIO()
                        with pd.ExcelWriter(output, engine='openpyxl') as writer:
                            df_final.to_excel(writer, index=False, sheet_name='Inventario Geral')
                            ws = writer.sheets['Inventario Geral']
                            for col in ws.columns:
                                max_len = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
                                ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4
                                
                                if col[0].value == "Peso":
                                    for cell in col[1:]:
                                        if cell.value is not None:
                                            cell.number_format = '0.000'

                        output.seek(0)
                        nome_download = f"{nome_arquivo_usuario.strip()}.xlsx" if nome_arquivo_usuario.strip() else "Inventario.xlsx"

                        st.success(f"✅ Sucesso! {len(todos_dfs)} arquivos processados.")
                        st.download_button(
                            label="📥 Baixar Excel Consolidado",
                            data=output,
                            file_name=nome_download,
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        )
                    else:
                        st.warning("Os arquivos foram lidos, mas nenhum dado válido foi encontrado.")

                except Exception as e:
                    st.error(f"Ocorreu um erro crítico: {e}")

# =====================================================================================
# ABA 2: LEITOR DE MÃO
# =====================================================================================
with aba_mao:
    if 'uploader_key_mao' not in st.session_state:
        st.session_state.uploader_key_mao = 0

    def limpar_lista_mao():
        st.session_state.uploader_key_mao += 1
        st.session_state.pop("resultado_mao", None)

    with st.expander("ℹ️ Como usar o Leitor de Mão", expanded=False):
        st.markdown(
            "1. Abra uma planilha Excel em branco e clique na célula **A1**.\n"
            "2. Ao chegar em cada local, **leia primeiro o QR da localização** (ex: `LOC:N4F1`) "
            "e depois todas as etiquetas daquele local.\n"
            "3. Ao terminar, salve a planilha e anexe abaixo.\n\n"
            "Formato esperado da etiqueta: `FILIAL-PRODUTO -ARMAZÉM-LOTE-PESO`\n\n"
            "🎨 **Legenda das cores:** 🟨 lote lido mais de uma vez · "
            "🟧 item sem localização · 🟥 leitura fora do padrão"
        )

    col1_m, col2_m = st.columns([2, 1])

    with col1_m:
        arquivos_mao = st.file_uploader(
            "Importar planilhas do leitor (.xlsx)",
            type="xlsx",
            accept_multiple_files=True,
            key=f"mao_{st.session_state.uploader_key_mao}",
            help="Planilhas Excel geradas pelo leitor de mão. Pode anexar várias de uma vez. "
                 "Cada planilha começa sem localização: leia o QR do local antes dos itens."
        )
        comparar_estoque = st.checkbox(
            "🔍 Comparar com o estoque do Protheus", value=False, key="comparar_est",
            help="Cruza os lotes lidos com o estoque do Protheus (aba Dados_Estoque, atualizada pelo robô) "
                 "da mesma filial das etiquetas, todos os armazéns. Mostra o que está no sistema e não "
                 "foi lido (Falta) e o que foi lido e não está no sistema (Sobra).")
        st.caption("Deixe desmarcado quando for só uma releitura de área para atualizar localização.")
        atualizar_base = st.checkbox(
            "📍 Atualizar Base de Localização", value=False, key="atualizar_base",
            help="Grava a localização dos itens lidos na Base de Localização online. É ela que a aba "
                 "Ordem de Carregamento consulta. Itens sem localização e leituras com erro não são gravados.")
        if arquivos_mao:
            col_msg_m, col_btn_m = st.columns([3, 1])
            with col_msg_m:
                st.info(f"📂 **{len(arquivos_mao)} arquivo(s) anexado(s).**")
            with col_btn_m:
                st.button("🗑️ Limpar Lista", on_click=limpar_lista_mao, type="secondary",
                          use_container_width=True, key="limpar_mao")

    with col2_m:
        st.info("Configurações de Saída")
        nome_arquivo_mao = st.text_input(
            "Nome do Arquivo Final (sem .xlsx):", value="", key="nome_mao",
            help='Se ficar em branco, o arquivo será salvo como "Inventario_LeitorMao.xlsx".'
        )
        data_inventario = st.date_input(
            "Data do inventário:", value=agora_br().date(), format="DD/MM/YYYY", key="data_inv",
            help="Dia em que as leituras foram feitas. Fica gravada no Excel e aparece no Roteiro de "
                 "Separação da aba Ordem de Carregamento."
        )

    st.markdown("###")
    if st.button("Converter Planilhas", type="primary", key="converter_mao"):
        if not arquivos_mao:
            st.warning("⚠️ Por favor, carregue pelo menos uma planilha .xlsx.")
        else:
            with st.spinner("Processando..."):
                try:
                    dfs = []
                    for arq in arquivos_mao:
                        try:
                            df_arq = processar_planilha_leitor(arq)
                            if df_arq.empty:
                                st.warning(f"⚠️ O arquivo **{arq.name}** não tem nenhuma leitura.")
                            else:
                                dfs.append(df_arq)
                        except Exception as e:
                            st.error(f"Erro ao ler o arquivo {arq.name}: {e}")

                    if dfs:
                        df_mao = marcar_duplicados(pd.concat(dfs, ignore_index=True))
                        df_locais = resumo_por_local(df_mao)

                        # Descrição do produto (lista vinda do Protheus via robô)
                        aviso_desc = ""
                        try:
                            produtos = carregar_produtos()
                        except Exception:
                            produtos = None
                            aviso_desc = ("Não consegui ler a lista de produtos do Protheus agora. "
                                          "O arquivo foi gerado com a coluna Descrição em branco.")
                        df_mao.insert(2, "Descrição",
                                      df_mao["Código"].map(produtos).fillna("") if produtos else "")
                        if produtos:
                            sem_desc = sorted(set(df_mao.loc[(df_mao["Código"] != "") &
                                                             (df_mao["Descrição"] == ""), "Código"]))
                            if sem_desc:
                                aviso_desc = (f"{len(sem_desc)} código(s) sem descrição (produto novo ou "
                                              f"código lido errado): {', '.join(sem_desc[:10])}"
                                              + (" ..." if len(sem_desc) > 10 else ""))
                        # Comparação com o estoque (opcional)
                        comparacao, aviso_estoque, hora_estoque = None, "", ""
                        if comparar_estoque:
                            try:
                                estoque, hora_estoque = carregar_estoque()
                                comparacao = comparar_com_estoque(df_mao, estoque)
                            except Exception:
                                aviso_estoque = ("Não consegui ler o estoque do Protheus agora. "
                                                 "A conversão foi feita sem a comparação.")

                        # Base de Localização (opcional)
                        msg_base, aviso_base = "", ""
                        if atualizar_base:
                            try:
                                origem = ", ".join(a.name for a in arquivos_mao)[:100]
                                n = gravar_base_localizacao(linhas_para_base(df_mao, origem))
                                msg_base = (f"📍 {n} leitura(s) gravada(s) na Base de Localização."
                                            if n else "📍 Nenhum item com localização para gravar na base.")
                            except Exception:
                                aviso_base = ("Não consegui gravar na Base de Localização agora. "
                                              "Tente converter de novo daqui a pouco.")

                        nome = nome_arquivo_mao.strip() or "Inventario_LeitorMao"
                        st.session_state.resultado_mao = {
                            "df": df_mao,
                            "locais": df_locais,
                            "excel": gerar_excel_leitor(df_mao, df_locais,
                                                        data_inventario.strftime("%d/%m/%Y"),
                                                        comparacao).getvalue(),
                            "nome": f"{nome}.xlsx",
                            "qtd_arquivos": len(dfs),
                            "aviso_desc": aviso_desc,
                            "comparacao": comparacao,
                            "aviso_estoque": aviso_estoque,
                            "hora_estoque": hora_estoque,
                            "msg_base": msg_base,
                            "aviso_base": aviso_base,
                        }
                    else:
                        st.session_state.pop("resultado_mao", None)
                except Exception as e:
                    st.error(f"Ocorreu um erro crítico: {e}")

    # --- RESULTADO (fica na tela mesmo depois de clicar em baixar) ---
    res = st.session_state.get("resultado_mao")
    if res:
        df_mao = res["df"]
        st.success(f"✅ Conversão concluída! {res['qtd_arquivos']} arquivo(s) processado(s).")
        if res.get("aviso_desc"):
            st.warning(f"⚠️ {res['aviso_desc']}")
        if res.get("msg_base"):
            st.info(res["msg_base"])
        if res.get("aviso_base"):
            st.warning(f"⚠️ {res['aviso_base']}")

        st.subheader("📊 Resumo da conversão")
        m1, m2, m3, m4, m5 = st.columns(5)
        m1.metric("Itens lidos", len(df_mao),
                  help="Total de leituras de etiquetas (sem contar os QR de localização).")
        m2.metric("Localizações", int((res["locais"]["Localização"] != "SEM LOCALIZAÇÃO").sum()),
                  help="Quantidade de locais diferentes lidos (QR com LOC:).")
        m3.metric("Lotes duplicados", int(df_mao.loc[df_mao["_dup"], "Lote"].nunique()),
                  help="Lotes lidos mais de uma vez. Todas as leituras repetidas ficam em amarelo.")
        m4.metric("Sem localização", int((df_mao["_status"] == "semloc").sum()),
                  help="Itens lidos antes de qualquer QR de localização na planilha. Ficam em laranja.")
        m5.metric("Leituras com erro", int((df_mao["_status"] == "erro").sum()),
                  help="Leituras fora do padrão FILIAL-PRODUTO -ARMAZÉM-LOTE-PESO. Ficam em vermelho.")

        st.markdown("#### 📍 Itens por localização")
        st.info("💡 Confira se as quantidades fazem sentido. Um local com muito mais itens que o normal, "
                "ou um local que deveria aparecer e não está na lista, pode indicar que alguém "
                "esqueceu de bipar a placa de localização.")
        df_loc = res["locais"]
        st.dataframe(
            df_loc.style.apply(
                lambda r: ["background-color: #FFE8CC" if r["Localização"] == "SEM LOCALIZAÇÃO" else ""] * len(r),
                axis=1),
            column_config={"Itens": st.column_config.ProgressColumn(
                "Itens", format="%d", min_value=0,
                max_value=int(df_loc["Itens"].max()) if not df_loc.empty else 1)},
            hide_index=True, use_container_width=False, width=520,
        )
        st.caption(f"Total: {int(df_loc['Itens'].sum())} itens. As leituras com erro não entram nesta "
                   "tabela. Os locais aparecem na ordem em que foram lidos.")

        st.markdown("#### 👀 Prévia do arquivo")
        st.caption("🟨 Lote lido mais de uma vez · 🟧 Sem localização · 🟥 Leitura fora do padrão")
        colunas_vis = ["Filial", "Código", "Descrição", "Armazém", "Lote", "Peso", "Localização", "Observação"]

        def _colorir(row):
            cor = cor_da_linha(df_mao.at[row.name, "_status"], df_mao.at[row.name, "_dup"])
            return [f"background-color: #{cor}" if cor else ""] * len(row)

        df_vis = df_mao[colunas_vis].copy()
        df_vis["Peso"] = df_vis["Peso"].apply(lambda v: "" if pd.isna(v) else f"{v:.3f}".replace(".", ","))
        st.dataframe(df_vis.style.apply(_colorir, axis=1), hide_index=True, use_container_width=True)
        st.caption("O Excel baixado sai com as mesmas cores, a coluna Observação e uma segunda aba "
                   "com os itens por localização.")

        # --- COMPARAÇÃO COM O ESTOQUE ---
        if res.get("aviso_estoque"):
            st.warning(f"⚠️ {res['aviso_estoque']}")
        comp = res.get("comparacao")
        if comp:
            st.subheader("🔍 Comparação com o estoque do Protheus")
            hora_txt = (f" · Estoque do Protheus copiado pelo robô em **{res['hora_estoque']}**"
                        if res.get("hora_estoque") else "")
            st.info(f"Filial **{', '.join(comp['filiais']) or '—'}**, todos os armazéns{hora_txt}. "
                    "Movimentações feitas durante o inventário (faturamento, produção, transferências) "
                    "podem gerar divergências que não são erro.")
            c1, c2, c3, c4, c5, c6 = st.columns(6)
            c1.metric("Lotes no Protheus", comp["n_protheus"],
                      help="Lotes com saldo no Protheus na filial das etiquetas.")
            c2.metric("Lotes lidos", comp["n_lidos"], help="Lotes diferentes lidos no inventário.")
            c3.metric("Encontrados", comp["n_encontrados"],
                      help="Lotes que estão no Protheus e foram lidos.")
            c4.metric("Falta", comp["n_falta"],
                      help="Lotes com saldo no Protheus que não foram lidos no inventário.")
            c5.metric("Sobra", comp["n_sobra"],
                      help="Lotes lidos que não estão no estoque do Protheus "
                           "(material sem entrada, já faturado ou etiqueta errada).")
            acur = comp["acuracidade"]
            c6.metric("Acuracidade", f"{acur:.1f}%".replace(".", ",") if acur is not None else "—",
                      help="Encontrados ÷ Lotes no Protheus.")

            aba_falta, aba_sobra = st.tabs([f"🟥 Falta ({comp['n_falta']})", f"🟦 Sobra ({comp['n_sobra']})"])
            with aba_falta:
                st.caption("Lotes com saldo no Protheus que **não foram lidos** no inventário.")
                st.dataframe(comp["falta"].style.format({"Saldo Protheus (t)": formatar_ton}),
                             hide_index=True, use_container_width=True)
            with aba_sobra:
                st.caption("Lotes **lidos** que não estão no estoque do Protheus, com o local onde foram lidos.")
                st.dataframe(comp["sobra"].style.format({"Peso lido (t)": formatar_ton}),
                             hide_index=True, use_container_width=True)
            st.caption("O Excel baixado ganha as abas **Falta** e **Sobra**.")

        st.download_button(
            label="📥 Baixar Excel Consolidado",
            data=res["excel"],
            file_name=res["nome"],
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="baixar_mao",
        )

# =====================================================================================
# ABA 3: ETIQUETAS DE LOCALIZAÇÃO
# =====================================================================================
with aba_etiquetas:
    st.info("Gere aqui os QR Codes das localizações para imprimir e fixar nos locais. "
            "Cada QR já sai com o prefixo `LOC:`, que o sistema usa para reconhecer que a "
            "leitura é um local e não um item.")

    col1_e, col2_e = st.columns([2, 1])
    with col1_e:
        texto_locais = st.text_area(
            "Localizações (uma por linha)", height=180, key="lista_locais",
            placeholder="N4F1\nN4F2\nN4F3",
            help="Digite ou cole as localizações, uma por linha. Pode colar direto de uma coluna "
                 "do Excel. Linhas em branco e nomes repetidos são ignorados."
        )
    with col2_e:
        layout_escolhido = st.selectbox(
            "Etiquetas por folha A4", list(LAYOUTS_ETIQUETA.keys()), index=2, key="layout_etq",
            help="Quanto menos etiquetas por folha, maior o QR. Etiquetas maiores são mais fáceis "
                 "de ler à distância."
        )
        nome_pdf = st.text_input("Nome do PDF (sem .pdf)", value="", key="nome_pdf",
                                 help='Se ficar em branco, será salvo como "Etiquetas_Localizacao.pdf".')

    if st.button("Gerar Etiquetas", type="primary", key="gerar_etq"):
        locais = limpar_lista_locais(texto_locais)
        if not locais:
            st.warning("⚠️ Digite pelo menos uma localização.")
            st.session_state.pop("resultado_etq", None)
        else:
            with st.spinner("Gerando etiquetas..."):
                cols_l, lins_l = LAYOUTS_ETIQUETA[layout_escolhido]
                paginas = -(-len(locais) // (cols_l * lins_l))
                st.session_state.resultado_etq = {
                    "locais": locais,
                    "pdf": gerar_pdf_etiquetas(locais, cols_l, lins_l).getvalue(),
                    "nome": f"{nome_pdf.strip() or 'Etiquetas_Localizacao'}.pdf",
                    "paginas": paginas,
                }

    res_e = st.session_state.get("resultado_etq")
    if res_e:
        st.success(f"✅ {len(res_e['locais'])} etiqueta(s) gerada(s) em {res_e['paginas']} página(s).")
        st.download_button("📥 Baixar PDF para Impressão", data=res_e["pdf"],
                           file_name=res_e["nome"], mime="application/pdf", key="baixar_etq")

        st.markdown("#### 👀 Prévia das etiquetas")
        LIMITE_PREVIA = 8
        mostrar = res_e["locais"][:LIMITE_PREVIA]
        cols_prev = st.columns(4)
        for i, local in enumerate(mostrar):
            with cols_prev[i % 4]:
                st.image(gerar_qr_png(PREFIXO_LOC + local, escala=6), width=150)
                st.markdown(f"<div style='font-size:24px;font-weight:700;margin-top:-8px'>{local}</div>",
                            unsafe_allow_html=True)
        if len(res_e["locais"]) > LIMITE_PREVIA:
            st.caption(f"Mostrando {LIMITE_PREVIA} de {len(res_e['locais'])}. O PDF contém todas.") 

# =====================================================================================
# ABA 4: ORDEM DE CARREGAMENTO
# =====================================================================================
with aba_ordem:
    if 'uploader_key_ordem' not in st.session_state:
        st.session_state.uploader_key_ordem = 0

    def limpar_lista_ordem():
        st.session_state.uploader_key_ordem += 1
        st.session_state.pop("resultado_ordem", None)

    with st.expander("ℹ️ Como usar a Ordem de Carregamento", expanded=False):
        st.markdown(
            "1. Anexe um ou mais PDFs de **Pick-list de Carregamento** gerados no Protheus.\n"
            "2. Escolha de onde vêm as localizações: da **Base de Localização online** (padrão, alimentada "
            "pela aba Leitor de Mão) ou de um **Excel do inventário** anexado.\n"
            "3. Clique em **Preencher Ordens** e baixe o PDF.\n\n"
            "📄 O PDF final traz, para cada carga: a **ordem original com a localização preenchida** "
            "na coluna *Localizac* e logo depois o **Roteiro de Separação**, com os itens agrupados por local.\n\n"
            "✏️ Lotes que não estão no inventário ficam com a localização **em branco**, para o "
            "conferente preencher à mão. No roteiro eles aparecem no final, em laranja."
        )

    origem_loc = st.radio(
        "Origem das localizações:", ["🌐 Base de Localização (online)", "📎 Anexar Excel do inventário"],
        horizontal=True, key="origem_loc",
        help="Base online: usa a leitura mais recente de cada lote gravada pela aba Leitor de Mão "
             "(caixinha 'Atualizar Base de Localização'). Excel: usa só o arquivo que você anexar.")
    usar_online = origem_loc.startswith("🌐")

    col1_o, col2_o = st.columns([1, 1])
    with col1_o:
        ordens_pdf = st.file_uploader(
            "1️⃣ Ordens de Carregamento (.pdf)", type="pdf", accept_multiple_files=True,
            key=f"ordens_{st.session_state.uploader_key_ordem}",
            help="PDFs 'Pick-list de Carregamento' do Protheus. Pode anexar várias ordens de uma vez."
        )
    with col2_o:
        if usar_online:
            bases_inv = []
            st.info("🌐 As localizações virão da **Base de Localização online**. No roteiro, a coluna "
                    "**Lido em** mostra há quantos dias cada lote foi lido.")
        else:
            bases_inv = st.file_uploader(
                "2️⃣ Base do inventário (.xlsx)", type="xlsx", accept_multiple_files=True,
                key=f"base_{st.session_state.uploader_key_ordem}",
                help="Excel convertido nas abas Leitor de Mão ou App Celular (aba 'Inventario Geral'). "
                     "Use sempre o inventário mais recente."
            )

    col_nome_o, col_data_o = st.columns([1, 1])
    with col_nome_o:
        nome_pdf_ordem = st.text_input(
            "Nome do PDF final (sem .pdf):", value="", key="nome_ordem",
            help='Se ficar em branco, será salvo como "Ordens_Preenchidas.pdf".')
    data_manual = ""
    with col_data_o:
        if not usar_online:
            data_manual = st.text_input(
                "Data do inventário (só se não for preenchida sozinha):", value="", key="data_manual",
                placeholder="DD/MM/AAAA",
                help="Excel gerado pela aba Leitor de Mão já traz a data gravada. Para arquivos antigos "
                     "ou do App Celular, digite aqui a data em que o inventário foi feito.")

    if ordens_pdf or bases_inv:
        col_msg_o, col_btn_o = st.columns([3, 1])
        with col_msg_o:
            st.info(f"📂 **{len(ordens_pdf or [])} ordem(ns)** e **{len(bases_inv or [])} base(s)** anexadas.")
        with col_btn_o:
            st.button("🗑️ Limpar Lista", on_click=limpar_lista_ordem, type="secondary",
                      use_container_width=True, key="limpar_ordem")

    st.markdown("###")
    if st.button("Preencher Ordens", type="primary", key="preencher_ordem"):
        if not ordens_pdf:
            st.warning("⚠️ Anexe pelo menos uma Ordem de Carregamento (.pdf).")
        elif not usar_online and not bases_inv:
            st.warning("⚠️ Anexe o Excel do inventário para buscar as localizações.")
        else:
            with st.spinner("Lendo ordens e buscando localizações..."):
                try:
                    from pypdf import PdfWriter
                    if usar_online:
                        base, lidos = montar_base(carregar_base_localizacao())
                        origem_txt = ("Base de Localização (consultada em "
                                      f"{agora_br().strftime('%d/%m/%Y %H:%M')})")
                    else:
                        base, datas = ler_base_inventario(bases_inv)
                        lidos = None
                        # Data que vai no roteiro
                        datas = sorted(set(d for d in datas if d))
                        if data_manual.strip():
                            data_base = data_manual.strip()
                        elif len(datas) == 1:
                            data_base = datas[0]
                        elif len(datas) > 1:
                            data_base = " / ".join(datas)
                        else:
                            data_base = "data não informada"
                        origem_txt = f"inventário de {data_base}"

                    pdf_final = PdfWriter()
                    resumo = []
                    for arq in ordens_pdf:
                        pdf_bytes = arq.getvalue()
                        cab, itens = ler_ordem_carregamento(io.BytesIO(pdf_bytes))
                        if not itens:
                            st.warning(f"⚠️ Não encontrei itens no arquivo **{arq.name}**. "
                                       "Confira se é um Pick-list de Carregamento do Protheus.")
                            continue
                        for pagina in carimbar_ordem(io.BytesIO(pdf_bytes), itens, base).pages:
                            pdf_final.add_page(pagina)
                        for pagina in gerar_roteiro(cab, itens, base, origem_txt, lidos).pages:
                            pdf_final.add_page(pagina)

                        n_ok = sum(1 for it in itens if buscar_local(base, it["lote"]))
                        clientes = list(dict.fromkeys(it["cliente"] for it in itens if it["cliente"]))
                        resumo.append({
                            "Carga": cab.get("carga", ""), "Pré-carga": cab.get("precarga", ""),
                            "Cliente(s)": ", ".join(clientes), "Itens": len(itens),
                            "Localizados": n_ok, "Sem localização": len(itens) - n_ok,
                        })

                    if resumo:
                        buf_pdf = io.BytesIO()
                        pdf_final.write(buf_pdf)
                        st.session_state.resultado_ordem = {
                            "pdf": buf_pdf.getvalue(),
                            "nome": f"{nome_pdf_ordem.strip() or 'Ordens_Preenchidas'}.pdf",
                            "resumo": pd.DataFrame(resumo),
                            "origem_txt": origem_txt,
                            "lotes_base": len(base),
                        }
                    else:
                        st.session_state.pop("resultado_ordem", None)
                except Exception as e:
                    st.error(f"Ocorreu um erro: {e}")

    # --- RESULTADO ---
    res_o = st.session_state.get("resultado_ordem")
    if res_o:
        df_res = res_o["resumo"]
        st.success(f"✅ {len(df_res)} ordem(ns) preenchida(s). Base com {res_o['lotes_base']} lotes "
                   f"({res_o['origem_txt']}).")

        st.subheader("📊 Resumo")
        r1, r2, r3, r4 = st.columns(4)
        r1.metric("Cargas", len(df_res), help="Quantidade de ordens de carregamento lidas.")
        r2.metric("Itens", int(df_res["Itens"].sum()), help="Total de linhas (lotes) somando todas as ordens.")
        r3.metric("Localizados", int(df_res["Localizados"].sum()),
                  help="Lotes encontrados no inventário. A localização foi escrita na ordem.")
        r4.metric("Sem localização", int(df_res["Sem localização"].sum()),
                  help="Lotes que não estão no inventário. Ficam em branco para o conferente preencher.")

        st.dataframe(
            df_res.style.apply(
                lambda r: ["background-color: #FFE8CC" if r["Sem localização"] > 0 else ""] * len(r), axis=1),
            hide_index=True, use_container_width=True)
        st.caption("🟧 Cargas com algum lote sem localização. O PDF traz, para cada carga, a ordem "
                   "preenchida seguida do Roteiro de Separação.")

        st.download_button("📥 Baixar PDF das Ordens", data=res_o["pdf"], file_name=res_o["nome"],
                           mime="application/pdf", key="baixar_ordem")                               